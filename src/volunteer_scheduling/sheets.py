from pathlib import Path
from typing import List, Union

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import schedule_2


def write_pretty_schedule_xlsx(
    grid: schedule_2.ScheduleGrid,
    volunteers: List[schedule_2.Volunteer],
    time_periods: List[schedule_2.TimePeriod],
    output_path: Union[str, Path] = "schedule.xlsx",
) -> None:
    """
    Write the schedule to a formatted XLSX spreadsheet.

    Layout:
        Column A: time period
        Column B: shift
        Columns C-I: Monday-Sunday

    Cells corresponding to unoperational shifts are painted red.
    """
    workbook = Workbook()
    worksheet = workbook.active
    if worksheet is None:
        raise ValueError("No active sheet in workbook")
    worksheet.title = "Weekly Schedule"
    worksheet.freeze_panes = "C2"

    schedule_grid = grid.s

    days = list(schedule_2.DayOfWeek)

    # ---------------------------------------------------------
    # Styles
    # ---------------------------------------------------------

    dark_blue_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    light_blue_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )
    day_off_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )
    unoperational_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC",
    )
    under_minimum_fill = PatternFill(
        fill_type="solid",
        fgColor="FCE5CD",
    )
    white_font = Font(
        color="FFFFFF",
        bold=True,
    )
    bold_font = Font(bold=True)

    thin_side = Side(
        style="thin",
        color="B7B7B7",
    )
    medium_side = Side(
        style="medium",
        color="666666",
    )

    normal_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    centered = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    left_aligned = Alignment(
        horizontal="left",
        vertical="top",
        wrap_text=True,
    )

    # ---------------------------------------------------------
    # Header
    # ---------------------------------------------------------

    worksheet.cell(row=1, column=1, value="Time Period")
    worksheet.cell(row=1, column=2, value="Shift")

    for column, day in enumerate(days, start=3):
        worksheet.cell(
            row=1,
            column=column,
            value=day.value,
        )

    for cell in worksheet[1]:
        cell.fill = dark_blue_fill
        cell.font = white_font
        cell.alignment = centered
        cell.border = normal_border

    worksheet.row_dimensions[1].height = 28

    # ---------------------------------------------------------
    # Day-off row
    # ---------------------------------------------------------

    day_off_row = 2

    worksheet.merge_cells(
        start_row=day_off_row,
        start_column=1,
        end_row=day_off_row,
        end_column=2,
    )

    day_off_cell = worksheet.cell(
        row=day_off_row,
        column=1,
        value="Day Off",
    )
    day_off_cell.font = bold_font
    day_off_cell.fill = day_off_fill
    day_off_cell.alignment = centered

    for column, day in enumerate(days, start=3):
        volunteers_off = sorted(
            volunteer.name
            for volunteer in volunteers
            if day in volunteer.days_off
        )

        cell = worksheet.cell(
            row=day_off_row,
            column=column,
            value=", ".join(volunteers_off),
        )
        cell.fill = day_off_fill
        cell.alignment = left_aligned

    for column in range(1, 10):
        worksheet.cell(
            row=day_off_row,
            column=column,
        ).border = normal_border

    worksheet.row_dimensions[day_off_row].height = 45

    # ---------------------------------------------------------
    # Find all shifts represented in the schedule
    # ---------------------------------------------------------

    all_shifts: List[schedule_2.Shift] = []
    seen_shifts = set()

    for day_schedule in schedule_grid.values():
        for shift_map in day_schedule.values():
            for shift in shift_map:
                if shift not in seen_shifts:
                    seen_shifts.add(shift)
                    all_shifts.append(shift)

    # Preserve the order of time_periods and group their shifts.
    shifts_by_period = {
        time_period: [
            shift
            for shift in all_shifts
            if shift.time_period == time_period
        ]
        for time_period in time_periods
    }

    # Include additional periods, such as a generated "Other" period.
    extra_periods = []

    for shift in all_shifts:
        if (
            shift.time_period not in shifts_by_period
            and shift.time_period not in extra_periods
        ):
            extra_periods.append(shift.time_period)

    ordered_periods = [*time_periods, *extra_periods]

    for time_period in extra_periods:
        shifts_by_period[time_period] = [
            shift
            for shift in all_shifts
            if shift.time_period == time_period
        ]

    # ---------------------------------------------------------
    # Time periods and shifts
    # ---------------------------------------------------------

    current_row = 3

    for time_period in ordered_periods:
        shifts = shifts_by_period.get(time_period, [])

        if not shifts:
            continue

        period_start_row = current_row
        period_end_row = current_row + len(shifts) - 1

        if period_start_row != period_end_row:
            worksheet.merge_cells(
                start_row=period_start_row,
                start_column=1,
                end_row=period_end_row,
                end_column=1,
            )

        period_cell = worksheet.cell(
            row=period_start_row,
            column=1,
            value=(
                f"{time_period.name}\n"
                f"({time_period.start:%H:%M} - "
                f"{time_period.end:%H:%M})"
            ),
        )
        period_cell.font = bold_font
        period_cell.fill = light_blue_fill
        period_cell.alignment = centered

        for shift in shifts:
            shift_cell = worksheet.cell(
                row=current_row,
                column=2,
                value=shift.name,
            )

            shift_cell.font = bold_font
            shift_cell.alignment = left_aligned

            for column, day in enumerate(days, start=3):
                cell = worksheet.cell(
                    row=current_row,
                    column=column,
                )
                if isinstance(cell, MergedCell):
                    raise ValueError("Unexpected MergedCell")

                cell.border = normal_border
                cell.alignment = left_aligned

                if day in shift.unoperational_days:
                    cell.value = "UNOPERATIONAL"
                    cell.fill = unoperational_fill
                    cell.font = Font(
                        color="9C0006",
                        bold=True,
                    )
                    cell.alignment = centered
                    continue

                assigned_volunteers = (
                    schedule_grid
                    .get(day, {})
                    .get(time_period, {})
                    .get(shift, [])
                )

                volunteer_names = sorted(
                    volunteer.name
                    for volunteer in assigned_volunteers
                )

                cell.value = ", ".join(volunteer_names)

                if len(assigned_volunteers) < shift.min_people:
                    cell.fill = under_minimum_fill

            worksheet.cell(
                row=current_row,
                column=1,
            ).border = normal_border
            shift_cell.border = normal_border

            worksheet.row_dimensions[current_row].height = 42
            current_row += 1

        # Add a stronger bottom border after each time period.
        for column in range(1, 10):
            cell = worksheet.cell(
                row=period_end_row,
                column=column,
            )
            cell.border = Border(
                left=cell.border.left,
                right=cell.border.right,
                top=cell.border.top,
                bottom=medium_side,
            )

    # ---------------------------------------------------------
    # Column dimensions and final formatting
    # ---------------------------------------------------------

    worksheet.column_dimensions["A"].width = 23
    worksheet.column_dimensions["B"].width = 24

    for column in range(3, 10):
        worksheet.column_dimensions[
            get_column_letter(column)
        ].width = 24

    worksheet.auto_filter.ref = f"A1:I{max(current_row - 1, 2)}"

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook.save(output_path)
    print(f"Schedule saved to {output_path}")
